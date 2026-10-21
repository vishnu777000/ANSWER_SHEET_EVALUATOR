import os
import tempfile
import unittest

from openpyxl import load_workbook

from modules.storage import save_to_excel


def _results(label="Q1"):
    return {
        "questions": [{"label": label, "marks": 5}],
        "total_marks": 5,
        "max_total": 5,
        "percentage": 100,
    }


class StorageTests(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.filepath = os.path.join(self.temp_dir.name, "results.xlsx")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_user_text_is_written_as_literal_and_readable(self):
        save_to_excel(
            self.filepath,
            "Databases",
            "=HYPERLINK(\"https://example.invalid\")",
            "@student",
            _results(),
        )

        workbook = load_workbook(self.filepath, data_only=False)
        try:
            sheet = workbook["Databases"]
            self.assertEqual(
                sheet.cell(row=2, column=1).value,
                "'=HYPERLINK(\"https://example.invalid\")",
            )
            self.assertNotEqual(sheet.cell(row=2, column=1).data_type, "f")
            self.assertEqual(sheet.cell(row=2, column=2).value, "'@student")
        finally:
            workbook.close()

    def test_different_question_layout_does_not_append_misaligned_marks(self):
        save_to_excel(
            self.filepath,
            "Databases",
            "Student One",
            "1",
            _results("Q1"),
        )

        with self.assertRaisesRegex(ValueError, "different question layout"):
            save_to_excel(
                self.filepath,
                "Databases",
                "Student Two",
                "2",
                _results("Q2"),
            )

        workbook = load_workbook(self.filepath, read_only=True)
        try:
            self.assertEqual(workbook["Databases"].max_row, 2)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
